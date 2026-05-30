"""测试 excel/generator.py 生成的 Excel 结构和公式"""

import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import openpyxl

from excel.generator import generate_excel


def _read_sheet(excel_bytes: io.BytesIO):
    excel_bytes.seek(0)
    wb = openpyxl.load_workbook(excel_bytes)
    return wb.active


def test_single_movie():
    movies = [{"name": "消失的人", "cumulative_share": 0.176, "show_count": 12000}]
    buf = generate_excel("6.15", movies, dapan_total=420000, total_show_count=360000)
    ws = _read_sheet(buf)

    # Row 1: headers
    assert ws.cell(1, 1).value == "日期"
    assert ws.cell(1, 10).value == "落位占比"

    # Row 2: data
    assert ws.cell(2, 1).value == "6.15"
    assert ws.cell(2, 2).value == "消失的人"
    assert ws.cell(2, 3).value == 0.176
    assert ws.cell(2, 4).value == 420000
    assert ws.cell(2, 5).value == 360000
    assert ws.cell(2, 8).value == 12000


def test_formulas():
    movies = [{"name": "A", "cumulative_share": 0.5, "show_count": 1000}]
    buf = generate_excel("6.15", movies, dapan_total=2000, total_show_count=1000)
    ws = _read_sheet(buf)

    assert ws.cell(2, 6).value == "=D2-E2"
    assert ws.cell(2, 7).value == "=F2*C2"
    assert ws.cell(2, 9).value == "=G2+H2"
    assert ws.cell(2, 10).value == "=ROUNDDOWN(I2/D2,3)"


def test_multi_row_formula_references():
    movies = [
        {"name": "A", "cumulative_share": 0.3, "show_count": 100},
        {"name": "B", "cumulative_share": 0.7, "show_count": 200},
    ]
    buf = generate_excel("6.15", movies, dapan_total=1000, total_show_count=500)
    ws = _read_sheet(buf)

    # Row 2 (first movie)
    assert ws.cell(2, 6).value == "=D2-E2"
    assert ws.cell(2, 9).value == "=G2+H2"

    # Row 3 (second movie) — row refs must shift
    assert ws.cell(3, 6).value == "=D3-E3"
    assert ws.cell(3, 7).value == "=F3*C3"
    assert ws.cell(3, 9).value == "=G3+H3"
    assert ws.cell(3, 10).value == "=ROUNDDOWN(I3/D3,3)"


def test_number_formats():
    movies = [{"name": "A", "cumulative_share": 0.5, "show_count": 1000}]
    buf = generate_excel("6.15", movies, dapan_total=2000, total_show_count=1000)
    ws = _read_sheet(buf)

    assert ws.cell(2, 3).number_format == "0.0%"   # cumulative_share
    assert ws.cell(2, 4).number_format == "#,##0"   # dapan
    assert ws.cell(2, 10).number_format == "0.0%"   # landing share


def test_empty_movies():
    buf = generate_excel("6.15", [], dapan_total=420000, total_show_count=360000)
    ws = _read_sheet(buf)

    # Only header row
    assert ws.cell(1, 1).value == "日期"
    assert ws.cell(2, 1).value is None
